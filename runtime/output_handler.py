import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx2pdf import convert

from runtime.providers.base import ResearchResult


REPO_ROOT = (
    Path(__file__)
    .resolve()
    .parents[1]
)


def safe_filename(value):
    value = str(value).strip()

    value = re.sub(
        r'[<>:"/\\|?*\n\r\t]+',
        "_",
        value
    )

    value = re.sub(
        r"\s+",
        "_",
        value
    )

    return value or "client"


def normalize_text(text):
    replacements = {
        "\u2010": "-",
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": " - ",
        "\u2212": "-",
        "\u00a0": " ",
        "\u202f": " ",
        "\u200b": "",
        "\ufeff": "",
    }

    for source, target in replacements.items():
        text = text.replace(
            source,
            target
        )

    return text


def get_output_directory(
    output_directory
):
    output_dir = (
        REPO_ROOT
        / output_directory
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    return output_dir


def save_markdown(
    report,
    filename,
    output_directory="output"
):
    output_dir = (
        get_output_directory(
            output_directory
        )
    )

    path = (
        output_dir
        / filename
    )

    path.write_text(
        report,
        encoding="utf-8"
    )

    return path


def append_source_appendix(
    report,
    sources
):
    """
    Append source metadata when citations
    are available but not rendered inline.
    """

    if not sources:
        return report

    existing_urls = set(
        re.findall(
            r"https?://\S+",
            report
        )
    )

    appendix = [
        "",
        "",
        "## Source Metadata Appendix",
        "",
        "The following sources were identified "
        "during the research process.",
        "",
    ]

    index = 1

    for source in sources:

        url = source.get(
            "url"
        )

        if not url:
            continue

        if url in existing_urls:
            continue

        title = (
            source.get("title")
            or "Untitled source"
        )

        appendix.append(
            f"{index}. {title} - {url}"
        )

        index += 1

    if index == 1:
        return report

    return (
        report
        + "\n"
        + "\n".join(
            appendix
        )
    )


def save_evidence_json(
    result,
    filename,
    output_directory="output"
):
    output_dir = (
        get_output_directory(
            output_directory
        )
    )

    path = (
        output_dir
        / filename
    )

    evidence = {
        "provider":
            result.provider,

        "citations":
            result.citations,

        "sources":
            result.sources,

        "metadata":
            result.metadata,
    }

    path.write_text(
        json.dumps(
            evidence,
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8"
    )

    return path


def set_cell_shading(
    cell,
    fill="E9EDF2"
):
    tc_pr = (
        cell._tc
        .get_or_add_tcPr()
    )

    shading = OxmlElement(
        "w:shd"
    )

    shading.set(
        qn("w:fill"),
        fill
    )

    tc_pr.append(
        shading
    )


def set_repeat_table_header(row):

    tr_pr = (
        row._tr
        .get_or_add_trPr()
    )

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true"
    )

    tr_pr.append(
        tbl_header
    )


def format_run(
    run,
    font_size=10.5,
    bold=False
):

    run.font.name = "Arial"

    run.font.size = Pt(
        font_size
    )

    run.bold = bold


def add_inline_text(
    paragraph,
    text,
    font_size=10.5
):

    parts = re.split(
        r"(\*\*.*?\*\*)",
        text
    )

    for part in parts:

        if not part:
            continue

        if (
            part.startswith("**")
            and part.endswith("**")
        ):

            run = paragraph.add_run(
                part[2:-2]
            )

            format_run(
                run,
                font_size,
                True
            )

        else:

            run = paragraph.add_run(
                part
            )

            format_run(
                run,
                font_size
            )


def configure_document(document):

    section = (
        document.sections[0]
    )

    section.top_margin = Inches(
        0.7
    )

    section.bottom_margin = Inches(
        0.7
    )

    section.left_margin = Inches(
        0.7
    )

    section.right_margin = Inches(
        0.7
    )

    normal = (
        document.styles[
            "Normal"
        ]
    )

    normal.font.name = "Arial"
    normal.font.size = Pt(
        10.5
    )


def is_table_separator(line):

    stripped = line.strip()

    if not (
        stripped.startswith("|")
        and stripped.endswith("|")
    ):
        return False

    cells = [
        x.strip()
        for x in (
            stripped
            .strip("|")
            .split("|")
        )
    ]

    for cell in cells:

        if not re.fullmatch(
            r":?-{3,}:?",
            cell
        ):
            return False

    return True


def parse_table_row(line):

    return [
        x.strip()
        for x in (
            line
            .strip()
            .strip("|")
            .split("|")
        )
    ]


def add_table(
    document,
    rows
):

    if not rows:
        return

    columns = max(
        len(row)
        for row in rows
    )

    table = document.add_table(
        rows=len(rows),
        cols=columns
    )

    table.style = (
        "Table Grid"
    )

    for row_index, row_data in enumerate(
        rows
    ):

        row = table.rows[
            row_index
        ]

        if row_index == 0:
            set_repeat_table_header(
                row
            )

        for col_index in range(
            columns
        ):

            cell = (
                row.cells[
                    col_index
                ]
            )

            text = (
                row_data[col_index]
                if col_index < len(row_data)
                else ""
            )

            paragraph = (
                cell.paragraphs[0]
            )

            add_inline_text(
                paragraph,
                text,
                8
            )

            if row_index == 0:

                set_cell_shading(
                    cell
                )

    document.add_paragraph()


def add_heading(
    document,
    text,
    level
):

    paragraph = document.add_heading(
        text,
        level=min(
            level,
            4
        )
    )

    return paragraph


def markdown_to_docx(
    report,
    docx_path
):

    document = Document()

    configure_document(
        document
    )

    lines = report.splitlines()

    index = 0

    while index < len(lines):

        line = (
            lines[index]
            .rstrip()
        )

        if not line.strip():

            index += 1
            continue


        heading = re.match(
            r"^(#{1,6})\s+(.+)$",
            line
        )

        if heading:

            level = len(
                heading.group(1)
            )

            add_heading(
                document,
                heading.group(2),
                level
            )

            index += 1
            continue


        if (
            line.strip().startswith("|")
            and index + 1 < len(lines)
            and is_table_separator(
                lines[index + 1]
            )
        ):

            rows = [
                parse_table_row(
                    line
                )
            ]

            index += 2

            while (
                index < len(lines)
                and lines[index]
                .strip()
                .startswith("|")
            ):

                rows.append(
                    parse_table_row(
                        lines[index]
                    )
                )

                index += 1


            add_table(
                document,
                rows
            )

            continue


        if line.startswith("- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            add_inline_text(
                paragraph,
                line[2:]
            )

            index += 1
            continue


        paragraph = document.add_paragraph()

        add_inline_text(
            paragraph,
            line
        )

        index += 1


    document.save(
        str(docx_path)
    )


def save_docx(
    report,
    filename,
    output_directory="output"
):

    output_dir = (
        get_output_directory(
            output_directory
        )
    )

    path = (
        output_dir
        / filename
    )

    markdown_to_docx(
        report,
        path
    )

    return path


def save_pdf_from_docx(
    docx_path,
    pdf_path
):

    try:

        convert(
            str(docx_path),
            str(pdf_path)
        )

    except Exception as exc:

        raise RuntimeError(
            f"PDF conversion failed: {exc}"
        )

    return pdf_path


def save_report(
    result,
    client_name,
    provider_name=None,
    output_directory="output"
):

    if not isinstance(
        result,
        ResearchResult
    ):

        raise TypeError(
            "Expected ResearchResult"
        )


    timestamp = (
        datetime.now()
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )


    client = safe_filename(
        client_name
    )

    provider = safe_filename(
        provider_name
        or result.provider
    )


    base_name = (
        "shell_screening_"
        f"{client}_"
        f"{provider}_"
        f"{timestamp}"
    )


    warnings = []


    final_report = append_source_appendix(
        result.text,
        result.sources
    )


    markdown_path = save_markdown(
        final_report,
        base_name + ".md",
        output_directory
    )


    evidence_path = None

    try:

        evidence_path = save_evidence_json(
            result,
            base_name
            + "_evidence.json",
            output_directory
        )

    except Exception as exc:

        warnings.append(
            f"Evidence save failed: {exc}"
        )


    docx_path = None

    try:

        docx_path = save_docx(
            final_report,
            base_name + ".docx",
            output_directory
        )

    except Exception as exc:

        warnings.append(
            f"DOCX generation failed: {exc}"
        )


    pdf_path = None


    if docx_path:

        try:

            output_dir = (
                get_output_directory(
                    output_directory
                )
            )

            pdf_path = save_pdf_from_docx(
                docx_path,
                output_dir
                /
                (
                    base_name
                    + ".pdf"
                )
            )


        except Exception as exc:

            warnings.append(
                f"PDF generation failed: {exc}"
            )


    return {

        "markdown":
            markdown_path,

        "evidence":
            evidence_path,

        "docx":
            docx_path,

        "pdf":
            pdf_path,

        "warnings":
            warnings,
    }